import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os


def validate_input():
    name = name_entry.get()
    score_text = score_entry.get()

    if not name or not score_text:
        messagebox.showerror("Input Error", "Please input data in all fields!")
        return None, None

    try:
        score = int(score_text)
    except ValueError:
        messagebox.showerror("Input Error", "Score must be a number only.")
        return None, None

    return name, score

def append_data_to_excel(name, score):
    filename = "student_scores.xlsx"
    if not os.path.exists(filename):
        wb = Workbook()
        ws = wb.active
        ws.title = "UserData"
        ws.append(["Name", "Score", "Remarks"])
        wb.save(filename)

    wb = load_workbook(filename)
    ws = wb["UserData"]

    for row in reversed(range(2, ws.max_row + 1)):
        if ws.cell(row=row, column=1).value == "Average Score:":
            ws.delete_rows(row - 1, 3)
            break

    remarks = "Passed" if score >= 75 else "Failed"
    ws.append([name, score, remarks])
    calculate_and_append_average(ws)
    wb.save(filename)
    wb.close()

def calculate_and_append_average(ws):
    scores = [
        row[1].value for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False)
        if isinstance(row[1].value, (int, float))
    ]
    if scores:
        avg_score = round(sum(scores) / len(scores), 2)
        remark = "Passed" if avg_score >= 75 else "Failed"

        ws.append([" ", " ", " "])
        ws.append(["Average Score:", avg_score, remark])

def view_data():
    try:
        wb = load_workbook("student_scores.xlsx")
        ws = wb["UserData"]
    except Exception as e:
        messagebox.showerror("Error", f"Failed to load data: {e}")
        return

    view_window = tk.Toplevel(window)
    view_window.title("Scores")

    for a, row in enumerate(ws.iter_rows(values_only=True)):
        for b, value in enumerate(row):
            tk.Label(view_window, text=str(value), padx=5, pady=2)\
                .grid(row=a, column=b, sticky="nsew", padx=1, pady=1)

def save_data():
    name, score = validate_input()
    if name is None or score is None:
        return

    append_data_to_excel(name, score)
    messagebox.showinfo("Done", "Data saved successfully.")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)


window = tk.Tk()
window.title("Students Score Tracker")
window.geometry("300x200")


tk.Label(window, text="Name:").grid(row=0, column=0, padx=10, pady=5, sticky="w")
tk.Label(window, text="Score:").grid(row=1, column=0, padx=10, pady=5, sticky="w")


name_entry = tk.Entry(window, width=30)
score_entry = tk.Entry(window, width=30)
name_entry.grid(row=0, column=1, pady=5)
score_entry.grid(row=1, column=1, pady=5)

tk.Button(window, text="Submit", command=save_data, width=20)\
    .grid(row=3, column=1, columnspan=2, pady=10)

tk.Button(window, text="View Scores", command=view_data, width=20)\
    .grid(row=4, column=1, columnspan=2, pady=5)

window.mainloop()
